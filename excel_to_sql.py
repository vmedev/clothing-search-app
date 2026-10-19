
from openpyxl import load_workbook
from pathlib import Path

# Excel faila atrašanās vieta
excel_file = "clothing_catalog_template.xlsx"

# Mape, kurā atrodas attēli
image_folder = Path("static/img")

# Izveidojam SQL komandas
workbook = load_workbook(excel_file, data_only=True)
sheet = workbook["Sheet1"]

sql_commands = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    product_id, image_name, color, brand, price, image_url, description = row

    # Izlaižam tukšās rindas
    if not product_id or not image_name:
        continue

    # Pārbauda, vai attēls eksistē
    if not (image_folder / str(image_name)).is_file():
        print(f"Brīdinājums: nav atrasts attēls {image_name}")

    # Droši sagatavojam tekstu SQL formātam
    def sql_text(value):
        if value is None:
            return "NULL"
        return "'" + str(value).replace("'", "''") + "'"

    price_sql = str(float(price)) if price is not None else "NULL"

    command = f"""
INSERT INTO PRODUCT
    (ID, IMAGE_NAME, COLOR, BRAND, PRICE, IMAGE_URL, DESCRIPTION)
VALUES
    ({int(product_id)}, {sql_text(image_name)}, {sql_text(color)},
     {sql_text(brand)}, {price_sql}, {sql_text(image_url)},
     {sql_text(description)});
""".strip()

    sql_commands.append(command)

# Saglabājam SQL failu
Path("data.sql").write_text(
    "\n\n".join(sql_commands) + "\n",
    encoding="utf-8"
)

print(f"data.sql izveidots! Ierakstu skaits: {len(sql_commands)}")
